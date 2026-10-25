import sys
sys.stdout.reconfigure(encoding='utf-8')
import os
import subprocess
import fitz
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# --- 1. GENERATE HTML RESUME ---
html_content = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Srinath Chowdary Amilineni - Resume (Goldman Sachs)</title>
<style>
  @page {
    size: A4;
    margin: 14mm 16mm 14mm 16mm;
  }

  body {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    color: #1a1a1a;
    line-height: 1.35;
    font-size: 9.5pt;
    margin: 0;
    padding: 0;
  }

  .header-running {
    font-size: 7.5pt;
    color: #777;
    margin-bottom: 8px;
    letter-spacing: 0.2px;
  }

  .name {
    font-family: "Georgia", "Times New Roman", serif;
    font-size: 21pt;
    font-weight: 500;
    color: #111;
    letter-spacing: 0.5px;
    margin: 0 0 3px 0;
  }

  .contact-info {
    font-size: 9.2pt;
    color: #333;
    line-height: 1.35;
    margin-bottom: 10px;
  }

  .contact-info a {
    color: #1a1a1a;
    text-decoration: none;
  }

  h2.section-title {
    font-family: "Georgia", "Times New Roman", serif;
    font-size: 12.5pt;
    font-weight: 600;
    color: #111;
    letter-spacing: 0.6px;
    text-transform: uppercase;
    margin-top: 10px;
    margin-bottom: 4px;
    padding-bottom: 1px;
    border-bottom: 1px solid #ccc;
  }

  .objective-text {
    font-size: 9.3pt;
    text-align: justify;
    margin: 0 0 6px 0;
  }

  .degree {
    font-weight: 600;
    font-size: 9.5pt;
  }

  .gpa {
    font-size: 9.3pt;
    margin-top: 1px;
    margin-bottom: 6px;
  }

  .skill-group {
    margin-bottom: 2.5px;
    font-size: 9.3pt;
  }

  .skill-label {
    font-weight: 600;
    color: #111;
  }

  .project {
    margin-bottom: 8px;
  }

  .project-header {
    margin-bottom: 2px;
  }

  .project-title {
    font-weight: 600;
    font-size: 9.8pt;
    color: #111;
  }

  .project-dates {
    font-weight: 400;
    color: #444;
    font-size: 9.2pt;
  }

  ul {
    margin: 2px 0 4px 0;
    padding-left: 17px;
  }

  li {
    margin-bottom: 2.5px;
    font-size: 9.2pt;
    text-align: justify;
  }

  .page-break {
    page-break-before: always;
  }

  .footer-running {
    display: flex;
    justify-content: space-between;
    font-size: 7.5pt;
    color: #777;
    margin-top: 15px;
    padding-top: 4px;
    border-top: 1px solid #eee;
  }

  @media print {
    .no-print { display: none; }
    body { margin: 0; }
  }
</style>
</head>
<body>

  <!-- PAGE 1 -->
  <div class="header-running">Srinath Chowdary Amilineni - Resume (Goldman Sachs)</div>

  <div class="name">AMILINENI SRINATH CHOWDARY</div>
  <div class="contact-info">
    Bangalore, Karnataka, India | +91 7799556852 | <a href="mailto:amilinenisrinath@gmail.com">amilinenisrinath@gmail.com</a><br>
    <a href="https://github.com/amilinenisrinath-lab">github.com/amilinenisrinath-lab</a> | <a href="https://linkedin.com/in/srinath-chowdary-amilineni">linkedin.com/in/srinath-chowdary-amilineni</a>
  </div>

  <h2 class="section-title">OBJECTIVE</h2>
  <div class="objective-text">
    Final-year B.Tech CSE (Data Science) undergraduate, graduating 2027, seeking a Software Engineering Analyst role at Goldman Sachs. Built and deployed data-driven web and cloud applications end to end in Python, TypeScript, and JavaScript-based stacks, with hands-on expertise in autonomous multi-agent AI systems, reliable financial data pipelines, and scalable backend architecture. Analytical problem solver who collaborates well with teammates and stakeholders and learns quickly from feedback.
  </div>

  <h2 class="section-title">EDUCATION</h2>
  <div class="degree">B.Tech, Computer Science Engineering (Data Science Specialization), 2023 - 2027</div>
  <div class="gpa">CGPA: 6.48 / 10</div>

  <h2 class="section-title">TECHNICAL SKILLS</h2>
  <div class="skill-group"><span class="skill-label">Programming Languages:</span> Python, TypeScript, Node.js, SQL, R</div>
  <div class="skill-group"><span class="skill-label">Web & Backend:</span> FastAPI, Next.js 14, React.js, Node.js, REST API design and integration, Server-Sent Events (SSE)</div>
  <div class="skill-group"><span class="skill-label">AI & Agentic Systems:</span> LangGraph (Multi-Agent Orchestration), Advanced Hybrid RAG, Qdrant Vector DB, Cross-Encoder Reranking, Ragas Evaluation Framework</div>
  <div class="skill-group"><span class="skill-label">Databases & Tools:</span> SQL, SQLite (SQLModel ORM), MongoDB (basics), Docker, Git, yfinance</div>
  <div class="skill-group"><span class="skill-label">Data & Analytics:</span> Exploratory data analysis, data cleaning & validation, trend analysis, demand forecasting, econometric visualization (R ggplot2, Recharts)</div>
  <div class="skill-group"><span class="skill-label">Engineering Practices:</span> Full software development lifecycle (requirements, design, implementation, testing, deployment, maintenance), modular and reusable code, collaborative team development</div>

  <h2 class="section-title">PROJECTS</h2>

  <div class="project">
    <div class="project-header">
      <span class="project-title">AlphaScribe AI: Multi-Agent Financial Research Analyst</span>, <span class="project-dates">Oct 2026 - Present</span>
    </div>
    <ul>
      <li>Architected an autonomous multi-agent research platform using <strong>LangGraph</strong> orchestrating 5 specialized nodes (Supervisor, Researcher, Analyst, Writer, Fact-Checker) to synthesize verified institutional investment memos from SEC 10-K filings.</li>
      <li>Engineered an <strong>Advanced Hybrid RAG pipeline</strong> combining dense cosine embeddings with sparse <strong>BM25 keyword search</strong> and <strong>cross-encoder reranking</strong> in <strong>Qdrant</strong>, achieving high-precision retrieval over complex multi-column balance sheets and financial tables.</li>
      <li>Eliminated LLM arithmetic hallucinations by integrating a sandboxed <strong>Python REPL</strong> tool for deterministic ratio calculations (CAGR, ROIC, debt service) and real-time market data ingestion via <strong>yfinance</strong>.</li>
      <li>Developed a polyglot visualization bridge executing <strong>R (ggplot2)</strong> scripts via subprocesses to generate publication-grade financial charts alongside interactive client-side <strong>Recharts</strong> visualizations.</li>
      <li>Constructed an event-driven full-stack system using <strong>FastAPI</strong> with <strong>Server-Sent Events (SSE)</strong> for real-time agent thought streaming, <strong>SQLite (SQLModel)</strong> for state persistence, and a modern <strong>Next.js 14 (TypeScript, Tailwind CSS)</strong> investor dashboard.</li>
      <li>Benchmarked system reliability quantitatively with the <strong>Ragas</strong> framework, achieving <strong>&gt;95% Faithfulness and Context Precision</strong> on benchmark SEC 10-K test cases.</li>
    </ul>
  </div>

  <div class="footer-running">
    <span>Srinath Chowdary Amilineni - Resume (Goldman Sachs)</span>
    <span>Page 1 of 2</span>
  </div>

  <!-- PAGE 2 -->
  <div class="page-break"></div>
  <div class="header-running">Srinath Chowdary Amilineni - Resume (Goldman Sachs)</div>

  <div class="project" style="margin-top: 10px;">
    <div class="project-header">
      <span class="project-title">Agricultural Data Analytics and Advisory System (Cloud Platform)</span>, <span class="project-dates">Jan 2026 - Apr 2026</span>
    </div>
    <ul>
      <li>Contributed across the full development cycle (design, coding, testing, deployment) of a cloud-based analytics and advisory application.</li>
      <li>Built multi-image upload, validation checks, and progress tracking; debugged data-integrity issues to keep inputs accurate and complete.</li>
      <li>Worked with the team on architecture decisions to keep modules maintainable, scalable, and easy to extend, and proposed design improvements as the system evolved.</li>
      <li>Implemented AI/ML-driven recommendation logic that turns analysis into actionable, client-facing suggestions.</li>
    </ul>
  </div>

  <div class="project">
    <div class="project-header">
      <span class="project-title">Canteen Management & Food Waste Reduction ("Renu")</span>, <span class="project-dates">Jul 2025 - Nov 2025</span>
    </div>
    <ul>
      <li>Designed and built a web-based analytics platform from scratch, taking it from requirements through to a deployed system with real-time dashboards.</li>
      <li>Engineered a structured data pipeline connecting a Smart Tray hardware device to backend processing and dashboard reporting.</li>
      <li>Developed modular, reusable demand-forecasting components to support data-driven decisions, and maintained the platform after release.</li>
    </ul>
  </div>

  <h2 class="section-title">CERTIFICATIONS & ACHIEVEMENTS</h2>
  <ul>
    <li><strong>IBM SkillsBuild:</strong> Data Analysis Fluency Pathways</li>
    <li><strong>Smart India Hackathon:</strong> Participant</li>
  </ul>

  <h2 class="section-title">STRENGTHS</h2>
  <ul>
    <li><strong>Clear written and spoken communication:</strong> explains technical findings and architecture decisions clearly to teammates and stakeholders</li>
    <li><strong>Analytical problem solving:</strong> structured, logical approach to complex systems; collaborative team player receptive to feedback</li>
  </ul>

  <div class="footer-running" style="margin-top: 40px;">
    <span>Srinath Chowdary Amilineni - Resume (Goldman Sachs)</span>
    <span>Page 2 of 2</span>
  </div>

</body>
</html>
"""

html_path = r"c:\Users\amili\OneDrive\문서\Project-Resume\docs\resume_updated.html"
with open(html_path, "w", encoding="utf-8") as f:
    f.write(html_content)

print(f"HTML resume written to {html_path}")

# --- 2. RENDER PDF USING EDGE HEADLESS ---
edge_exe = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
pdf_out_docs = r"c:\Users\amili\OneDrive\문서\Project-Resume\docs\Srinath_Chowdary_Amilineni_Resume_Goldman_Sachs.pdf"
pdf_out_downloads = r"C:\Users\amili\Downloads\Srinath_Chowdary_Amilineni_Resume_Goldman_Sachs.pdf"

cmd = [
    edge_exe,
    "--headless",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={pdf_out_docs}",
    html_path
]
subprocess.run(cmd, check=True)
print(f"PDF generated at {pdf_out_docs}")

import shutil
shutil.copy(pdf_out_docs, pdf_out_downloads)
print(f"PDF copied to {pdf_out_downloads}")

# Verify generated PDF
doc_test = fitz.open(pdf_out_docs)
print(f"Generated PDF page count: {len(doc_test)}")
for i, p in enumerate(doc_test):
    print(f"Page {i+1} character count: {len(p.get_text())}")

# --- 3. CREATE MICROSOFT WORD (.DOCX) RESUME ---
doc = docx.Document()

# Set page margins
sections = doc.sections
for s in sections:
    s.top_margin = Inches(0.7)
    s.bottom_margin = Inches(0.7)
    s.left_margin = Inches(0.75)
    s.right_margin = Inches(0.75)

# Helper function for setting font
def set_font(run, font_name="Calibri", size_pt=10, bold=False, color_rgb=(26, 26, 26)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)

def add_section_heading(document, title_text):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_text)
    set_font(r, font_name="Georgia", size_pt=13, bold=True, color_rgb=(17, 17, 17))
    
    # Bottom border on paragraph
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), 'CCCCCC')
    pBdr.append(bottom)
    pPr.append(pBdr)

# Header: Name
p_name = doc.add_paragraph()
p_name.paragraph_format.space_before = Pt(0)
p_name.paragraph_format.space_after = Pt(2)
r_name = p_name.add_run("AMILINENI SRINATH CHOWDARY")
set_font(r_name, font_name="Georgia", size_pt=20, bold=True, color_rgb=(17, 17, 17))

# Contact
p_contact = doc.add_paragraph()
p_contact.paragraph_format.space_after = Pt(2)
r_c1 = p_contact.add_run("Bangalore, Karnataka, India | +91 7799556852 | amilinenisrinath@gmail.com\n")
set_font(r_c1, font_name="Calibri", size_pt=9.5, color_rgb=(60, 60, 60))
r_c2 = p_contact.add_run("github.com/amilinenisrinath-lab | linkedin.com/in/srinath-chowdary-amilineni")
set_font(r_c2, font_name="Calibri", size_pt=9.5, color_rgb=(60, 60, 60))

# OBJECTIVE
add_section_heading(doc, "OBJECTIVE")
p_obj = doc.add_paragraph()
p_obj.paragraph_format.space_after = Pt(6)
p_obj.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r_obj = p_obj.add_run("Final-year B.Tech CSE (Data Science) undergraduate, graduating 2027, seeking a Software Engineering Analyst role at Goldman Sachs. Built and deployed data-driven web and cloud applications end to end in Python, TypeScript, and JavaScript-based stacks, with hands-on expertise in autonomous multi-agent AI systems, reliable financial data pipelines, and scalable backend architecture. Analytical problem solver who collaborates well with teammates and stakeholders and learns quickly from feedback.")
set_font(r_obj, font_name="Calibri", size_pt=9.5)

# EDUCATION
add_section_heading(doc, "EDUCATION")
p_edu = doc.add_paragraph()
p_edu.paragraph_format.space_after = Pt(1)
r_edu1 = p_edu.add_run("B.Tech, Computer Science Engineering (Data Science Specialization), 2023 - 2027")
set_font(r_edu1, font_name="Calibri", size_pt=10, bold=True)
p_cgpa = doc.add_paragraph()
p_cgpa.paragraph_format.space_after = Pt(6)
r_cgpa = p_cgpa.add_run("CGPA: 6.48 / 10")
set_font(r_cgpa, font_name="Calibri", size_pt=9.5)

# TECHNICAL SKILLS
add_section_heading(doc, "TECHNICAL SKILLS")
skills_data = [
    ("Programming Languages", "Python, TypeScript, Node.js, SQL, R"),
    ("Web & Backend", "FastAPI, Next.js 14, React.js, Node.js, REST API design and integration, Server-Sent Events (SSE)"),
    ("AI & Agentic Systems", "LangGraph (Multi-Agent Orchestration), Advanced Hybrid RAG, Qdrant Vector DB, Cross-Encoder Reranking, Ragas Evaluation Framework"),
    ("Databases & Tools", "SQL, SQLite (SQLModel ORM), MongoDB (basics), Docker, Git, yfinance"),
    ("Data & Analytics", "Exploratory data analysis, data cleaning & validation, trend analysis, demand forecasting, econometric visualization (R ggplot2, Recharts)"),
    ("Engineering Practices", "Full software development lifecycle (requirements, design, implementation, testing, deployment, maintenance), modular and reusable code, collaborative team development")
]
for label, val in skills_data:
    p_sk = doc.add_paragraph()
    p_sk.paragraph_format.space_after = Pt(2.5)
    r_l = p_sk.add_run(f"• {label}: ")
    set_font(r_l, font_name="Calibri", size_pt=9.5, bold=True)
    r_v = p_sk.add_run(val)
    set_font(r_v, font_name="Calibri", size_pt=9.5)

# PROJECTS
add_section_heading(doc, "PROJECTS")

# Project 1: AlphaScribe AI
p_p1 = doc.add_paragraph()
p_p1.paragraph_format.space_before = Pt(4)
p_p1.paragraph_format.space_after = Pt(2)
p_p1.paragraph_format.keep_with_next = True
r_p1_t = p_p1.add_run("AlphaScribe AI: Multi-Agent Financial Research Analyst")
set_font(r_p1_t, font_name="Calibri", size_pt=10, bold=True)
r_p1_d = p_p1.add_run(", Oct 2026 - Present")
set_font(r_p1_d, font_name="Calibri", size_pt=9.5)

p1_bullets = [
    ("Architected an autonomous multi-agent research platform", " using LangGraph orchestrating 5 specialized nodes (Supervisor, Researcher, Analyst, Writer, Fact-Checker) to synthesize verified institutional investment memos from SEC 10-K filings."),
    ("Engineered an Advanced Hybrid RAG pipeline", " combining dense cosine embeddings with sparse BM25 keyword search and cross-encoder reranking in Qdrant, achieving high-precision retrieval over complex multi-column balance sheets and financial tables."),
    ("Eliminated LLM arithmetic hallucinations", " by integrating a sandboxed Python REPL tool for deterministic ratio calculations (CAGR, ROIC, debt service) and real-time market data ingestion via yfinance."),
    ("Developed a polyglot visualization bridge", " executing R (ggplot2) scripts via subprocesses to generate publication-grade financial charts alongside interactive client-side Recharts visualizations."),
    ("Constructed an event-driven full-stack system", " using FastAPI with Server-Sent Events (SSE) for real-time agent thought streaming, SQLite (SQLModel) for state persistence, and a modern Next.js 14 (TypeScript, Tailwind CSS) investor dashboard."),
    ("Benchmarked system reliability quantitatively", " using the Ragas framework, achieving >95% Faithfulness and Context Precision on benchmark SEC 10-K test cases.")
]
for bold_prefix, rest in p1_bullets:
    p_b = doc.add_paragraph(style='List Bullet')
    p_b.paragraph_format.space_after = Pt(2.5)
    p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_b = p_b.add_run(bold_prefix)
    set_font(r_b, font_name="Calibri", size_pt=9.5, bold=True)
    r_r = p_b.add_run(rest)
    set_font(r_r, font_name="Calibri", size_pt=9.5)

# Project 2: Ag Analytics
p_p2 = doc.add_paragraph()
p_p2.paragraph_format.space_before = Pt(4)
p_p2.paragraph_format.space_after = Pt(2)
p_p2.paragraph_format.keep_with_next = True
r_p2_t = p_p2.add_run("Agricultural Data Analytics and Advisory System (Cloud Platform)")
set_font(r_p2_t, font_name="Calibri", size_pt=10, bold=True)
r_p2_d = p_p2.add_run(", Jan 2026 - Apr 2026")
set_font(r_p2_d, font_name="Calibri", size_pt=9.5)

p2_bullets = [
    "Contributed across the full development cycle (design, coding, testing, deployment) of a cloud-based analytics and advisory application.",
    "Built multi-image upload, validation checks, and progress tracking; debugged data-integrity issues to keep inputs accurate and complete.",
    "Worked with the team on architecture decisions to keep modules maintainable, scalable, and easy to extend, and proposed design improvements as the system evolved.",
    "Implemented AI/ML-driven recommendation logic that turns analysis into actionable, client-facing suggestions."
]
for b_text in p2_bullets:
    p_b = doc.add_paragraph(style='List Bullet')
    p_b.paragraph_format.space_after = Pt(2.5)
    p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_b = p_b.add_run(b_text)
    set_font(r_b, font_name="Calibri", size_pt=9.5)

# Project 3: Canteen Management
p_p3 = doc.add_paragraph()
p_p3.paragraph_format.space_before = Pt(4)
p_p3.paragraph_format.space_after = Pt(2)
p_p3.paragraph_format.keep_with_next = True
r_p3_t = p_p3.add_run("Canteen Management & Food Waste Reduction (\"Renu\")")
set_font(r_p3_t, font_name="Calibri", size_pt=10, bold=True)
r_p3_d = p_p3.add_run(", Jul 2025 - Nov 2025")
set_font(r_p3_d, font_name="Calibri", size_pt=9.5)

p3_bullets = [
    "Designed and built a web-based analytics platform from scratch, taking it from requirements through to a deployed system with real-time dashboards.",
    "Engineered a structured data pipeline connecting a Smart Tray hardware device to backend processing and dashboard reporting.",
    "Developed modular, reusable demand-forecasting components to support data-driven decisions, and maintained the platform after release."
]
for b_text in p3_bullets:
    p_b = doc.add_paragraph(style='List Bullet')
    p_b.paragraph_format.space_after = Pt(2.5)
    p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_b = p_b.add_run(b_text)
    set_font(r_b, font_name="Calibri", size_pt=9.5)

# CERTIFICATIONS & ACHIEVEMENTS
add_section_heading(doc, "CERTIFICATIONS & ACHIEVEMENTS")
certs = [
    ("IBM SkillsBuild:", " Data Analysis Fluency Pathways"),
    ("Smart India Hackathon:", " Participant")
]
for c_b, c_r in certs:
    p_c = doc.add_paragraph(style='List Bullet')
    p_c.paragraph_format.space_after = Pt(2)
    r1 = p_c.add_run(c_b)
    set_font(r1, font_name="Calibri", size_pt=9.5, bold=True)
    r2 = p_c.add_run(c_r)
    set_font(r2, font_name="Calibri", size_pt=9.5)

# STRENGTHS
add_section_heading(doc, "STRENGTHS")
strengths = [
    ("Clear written and spoken communication:", " explains technical findings and architecture decisions clearly to teammates and stakeholders"),
    ("Analytical problem solving:", " structured, logical approach to complex systems; collaborative team player receptive to feedback")
]
for s_b, s_r in strengths:
    p_s = doc.add_paragraph(style='List Bullet')
    p_s.paragraph_format.space_after = Pt(2)
    r1 = p_s.add_run(s_b)
    set_font(r1, font_name="Calibri", size_pt=9.5, bold=True)
    r2 = p_s.add_run(s_r)
    set_font(r2, font_name="Calibri", size_pt=9.5)

docx_out_docs = r"c:\Users\amili\OneDrive\문서\Project-Resume\docs\Srinath_Chowdary_Amilineni_Resume_Goldman_Sachs.docx"
docx_out_downloads = r"C:\Users\amili\Downloads\Srinath_Chowdary_Amilineni_Resume_Goldman_Sachs.docx"

doc.save(docx_out_docs)
shutil.copy(docx_out_docs, docx_out_downloads)
print(f"DOCX resume saved at {docx_out_docs} and copied to {docx_out_downloads}")
